import io,json,math,zipfile,posixpath,uuid,re
import xml.etree.ElementTree as ET
from datetime import datetime
import shared_store as shared
NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
HEADERS=['Product code','Product name','Category','Current quantity','New quantity','Current price (INR)','New price (INR)']
def export_data(connect):
 with connect() as db:
  rows=[]
  for r in db.execute('SELECT p.*,i.record FROM products p JOIN inventory i USING(product_code) WHERE active=1 AND published=1 ORDER BY position'):
   p=json.loads(r['record']);rows.append([r['product_code'],r['name'],r['category'],shared.stock(p),None,r['price'],None])
  return dict(rows=rows,revision=db.execute("SELECT value FROM shared_meta WHERE key='revision'").fetchone()[0])
def parse(raw):
 try:
  with zipfile.ZipFile(io.BytesIO(raw)) as z:
   if len(z.infolist())>500 or sum(x.file_size for x in z.infolist())>20000000:raise ValueError('Workbook is too large.')
   def xml(path):
    data=z.read(path)
    if b'<!DOCTYPE' in data or b'<!ENTITY' in data:raise ValueError('Unsupported workbook XML.')
    return ET.fromstring(data)
   strings=[]
   if 'xl/sharedStrings.xml' in z.namelist():strings=[''.join(x.itertext()) for x in xml('xl/sharedStrings.xml').findall('m:si',NS)]
   wb=xml('xl/workbook.xml'); sheet=next((s for s in wb.findall('m:sheets/m:sheet',NS) if s.attrib.get('name')=='Catalogue'),None)
   if sheet is None:raise ValueError('The Catalogue sheet is missing.')
   rid=sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
   rel=next(r for r in xml('xl/_rels/workbook.xml.rels') if r.attrib['Id']==rid)
   path=posixpath.normpath(posixpath.join('xl',rel.attrib['Target'])) if not rel.attrib['Target'].startswith('/') else rel.attrib['Target'].lstrip('/')
   cells={}
   for c in xml(path).findall('m:sheetData/m:row/m:c',NS):
    address=c.attrib['r'];v=c.find('m:v',NS);value=v.text if v is not None else None
    if c.find('m:f',NS) is not None and (address.startswith(('A','E','G')) or address=='B5'):raise ValueError('Enter values, not formulas, in product codes and update columns.')
    if c.attrib.get('t')=='s':value=strings[int(value)]
    elif c.attrib.get('t')=='inlineStr':value=''.join(c.find('m:is',NS).itertext())
    cells[address]=value
   if [cells.get(c+'7') for c in 'ABCDEFG']!=HEADERS:raise ValueError('Column headings changed. Download the catalogue template again.')
   revision=int(cells.get('B5',''))
   numbers=sorted({int(re.search(r'\d+$',k).group()) for k in cells if re.fullmatch(r'[AEG]\d+',k) and int(k[1:])>=8})
   seen=set();updates=[]
   for row in numbers:
    code=cells.get('A'+str(row));qty=cells.get('E'+str(row));price=cells.get('G'+str(row))
    if not code:
     if qty is not None or price is not None:raise ValueError('Product code missing on row '+str(row))
     continue
    if code in seen:raise ValueError('Duplicate product code: '+code)
    seen.add(code)
    update={'code':code}
    for key,value in [('quantity',qty),('price',price)]:
     if value is None or not str(value).strip():continue
     try:n=float(value)
     except (ValueError,TypeError):raise ValueError('Invalid '+key+' for '+code)
     if not math.isfinite(n) or n<0 or n>10000000 or (key=='quantity' and (n!=int(n) or n>1000000)) or (key=='price' and abs(n*100-round(n*100))>1e-6):raise ValueError('Invalid '+key+' for '+code)
     update[key]=int(n) if key=='quantity' else round(n,2)
    if len(update)>1:updates.append(update)
   return revision,updates,seen
 except (zipfile.BadZipFile,KeyError,ET.ParseError,StopIteration,TypeError,IndexError):raise ValueError('Invalid .xlsx file. Use the downloaded catalogue template.')
def import_sheet(connect,raw,apply=False):
 revision,updates,codes=parse(raw)
 with connect() as db:
  db.execute('BEGIN IMMEDIATE')
  if revision!=db.execute("SELECT value FROM shared_meta WHERE key='revision'").fetchone()[0]:raise shared.Conflict('The catalogue changed after this sheet was downloaded. Download a fresh sheet and copy your changes into it.')
  existing={r['product_code']:r for r in db.execute('SELECT * FROM products WHERE active=1 AND published=1')}
  unknown=codes-set(existing)
  if unknown:raise ValueError('Unknown product code: '+sorted(unknown)[0])
  changes=[]
  for u in updates:
   r=existing[u['code']];p=json.loads(db.execute('SELECT record FROM inventory WHERE product_code=?',(u['code'],)).fetchone()[0]);oldqty=shared.stock(p)
   qty=u.get('quantity',oldqty);price=u.get('price',r['price']);soldout=int(qty<=0) if 'quantity' in u else r['sold_out']
   if qty==oldqty and price==r['price'] and soldout==r['sold_out']:continue
   changes.append(dict(code=u['code'],name=r['name'],oldQuantity=oldqty,newQuantity=qty,oldPrice=r['price'],newPrice=price))
   if apply:
    delta=qty-oldqty;p['adjustment']=p.get('adjustment',0)+delta;p['sp']=price;p['upd']=datetime.now().date().isoformat()
    db.execute('UPDATE products SET price=?,sold_out=? WHERE product_code=?',(price,soldout,u['code']))
    db.execute('UPDATE inventory SET record=? WHERE product_code=?',(json.dumps(p),u['code']))
    if delta:
     t=dict(id='excel-'+uuid.uuid4().hex,date=p['upd'],type='Adjustment',code=u['code'],qty=delta,ref='Excel stock update',rate=None)
     db.execute('INSERT INTO inventory_transactions VALUES (?,?,?)',(t['id'],u['code'],json.dumps(t)))
  if apply and changes:db.execute("UPDATE shared_meta SET value=value+1 WHERE key='revision'")
  return dict(count=len(changes),changes=changes,applied=apply)


def export_workbook(connect):
 from openpyxl import Workbook
 from openpyxl.styles import Font, PatternFill, Alignment
 from openpyxl.worksheet.datavalidation import DataValidation
 from openpyxl.worksheet.table import Table, TableStyleInfo
 data=export_data(connect);wb=Workbook();ws=wb.active;ws.title='Catalogue'
 ws['A2']='Kalajay catalogue updates'
 ws['A2'].font=Font(size=16,bold=True,color='A4502B')
 ws['A3']='Fill yellow cells only. Blank = no change; new quantity 0 = sold out.'
 ws['A4']='Upload the edited workbook in ERP → Catalogue Excel. Download a fresh sheet before editing.'
 ws['A5']='Database version';ws['B5']=data['revision']
 ws['A6']='Source: shared Kalajay catalogue'
 for col,value in enumerate(HEADERS,1):ws.cell(7,col,value)
 for row in data['rows']:ws.append(row) if ws.max_row>=7 else None
 end=ws.max_row;ws.freeze_panes='A8'
 if end>=8:
  tab=Table(displayName='CatalogueUpdates',ref=f'A7:G{end}');tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);ws.add_table(tab)
  for col,kind in [('E','whole'),('G','decimal')]:
   validation=DataValidation(type=kind,operator='greaterThanOrEqual',formula1=0,allow_blank=True);validation.showErrorMessage=True;validation.error='Enter a non-negative value.';ws.add_data_validation(validation);validation.add(f'{col}8:{col}{end}')
  for row in ws.iter_rows(min_row=8):
   for i in (4,6):row[i].fill=PatternFill('solid',fgColor='FFF2CC')
   for i in (5,6):row[i].number_format='0.00'
 for col,width in [('A',20),('B',48),('C',24),('D',20),('E',20),('F',24),('G',24)]:ws.column_dimensions[col].width=width
 out=io.BytesIO();wb.save(out);return out.getvalue()
